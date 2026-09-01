"""Build the visual V2 AI GURU product thesis from Markdown.

The Markdown remains the source of truth for prose.  V2 adds a visual layer:
the diagrams are authored as SVG, wrapped as self-contained HTML for review,
and rasterized to crisp PNG assets for reliable Word rendering.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "AI_GURU_THUYET_MINH_SAN_PHAM_V2.md"
DEFAULT_OUTPUT = ROOT / "AI_GURU_THUYET_MINH_SAN_PHAM_V2.docx"

FONT = "Calibri"
MONO_FONT = "Consolas"
INK = "0B2545"
NAVY = "0B2545"
TEAL = "1F8A89"
DARK_TEAL = "176B6A"
CORAL = "E88C55"
CORAL_DARK = "D46F3D"
SLATE = "52657A"
MUTED = "7A8A99"
LIGHT = "F4F6F6"
LIGHT_TEAL = "E5F2EF"
LIGHT_CORAL = "FEF0E7"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F0F3F5"
PAPER = "FAF9F5"
WHITE = "FFFFFF"
RED = "9B1C1C"
GREEN = "1F3A5F"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_DXA = {"top": 90, "bottom": 90, "start": 120, "end": 120}

IMAGE_DIR = ROOT / "assets" / "v2_png"
IMAGE_CAPTIONS = {
    "cover_flow_v2": "Hình 1. Mạch sản phẩm: từ câu hỏi đến câu trả lời có thể kiểm chứng.",
    "data_normalization_v2": "Hình 2. Chuẩn hóa dữ liệu: làm dữ liệu dễ dùng hơn nhưng vẫn giữ đường quay về nguồn.",
    "table_semantics_v2": "Hình 3. Một exact cell chỉ có ý nghĩa khi đi cùng bảng, header, dòng, cột và đơn vị.",
    "architecture_v2": "Hình 4. Kiến trúc ba vùng của AI GURU và ranh giới giữa đề xuất với kiểm chứng.",
    "pipeline_v2": "Hình 5. Pipeline sản phẩm từ Intake đến Delivery và Feedback.",
    "status_ladder_v2": "Hình 6. Các trạng thái tin cậy: mỗi bậc cần thêm một lớp bằng chứng.",
    "example_v2": "Hình 7. Ví dụ tính phần trăm tăng trưởng có thể truy vết.",
    "run_guide_v2": "Hình 8. Bốn chặng vận hành từ chuẩn bị môi trường đến đọc output.",
}


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(
    run,
    *,
    name: str = FONT,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_paragraph_border(paragraph, *, color: str = TEAL, size: str = "12") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, margins: dict[str, int] | None = None) -> None:
    margins = margins or CELL_MARGIN_DXA
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in margins.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, *, color: str = "D7DEE8", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), size)
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Inches(width_dxa / 1440)


def set_table_geometry(table, widths: list[int], *, indent: int = TABLE_INDENT_DXA) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}: {widths}")

    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)
    set_run_font(run, size=9, color=MUTED)


def configure_styles(doc: Document) -> None:
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb("1F2937")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for style_name, size, color, before, after in (
        ("Heading 1", 17, NAVY, 20, 10),
        ("Heading 2", 13, TEAL, 14, 6),
        ("Heading 3", 12, DARK_TEAL, 9, 4),
    ):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True

    def add_style(name: str, base: str = "Normal"):
        if name in styles:
            return styles[name]
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles[base]
        return style

    body = add_style("Narrative Body")
    body.font.name = FONT
    body._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    body._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    body.font.size = Pt(11)
    body.font.color.rgb = rgb("1F2937")
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.space_after = Pt(8)
    body.paragraph_format.line_spacing = 1.333

    code = add_style("Code Block")
    code.font.name = MONO_FONT
    code._element.rPr.rFonts.set(qn("w:ascii"), MONO_FONT)
    code._element.rPr.rFonts.set(qn("w:hAnsi"), MONO_FONT)
    code.font.size = Pt(8.5)
    code.font.color.rgb = rgb(INK)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0

    caption = add_style("Figure Caption")
    caption.font.name = FONT
    caption._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    caption.font.size = Pt(9.5)
    caption.font.italic = True
    caption.font.color.rgb = rgb(MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)

    kicker = add_style("Chapter Kicker")
    kicker.font.name = FONT
    kicker._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    kicker._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    kicker.font.size = Pt(8.5)
    kicker.font.bold = True
    kicker.font.color.rgb = rgb(TEAL)
    kicker.paragraph_format.space_before = Pt(14)
    kicker.paragraph_format.space_after = Pt(2)
    kicker.paragraph_format.keep_with_next = True

    image_caption = add_style("Image Caption")
    image_caption.base_style = styles["Figure Caption"]


def configure_page_furniture(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    header_p = header.paragraphs[0]
    header_p.text = ""
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    run = header_p.add_run("AI GURU")
    set_run_font(run, size=8.5, color=TEAL, bold=True)
    run = header_p.add_run("  |  THUYẾT MINH SẢN PHẨM")
    set_run_font(run, size=8.5, color=SLATE, bold=True)
    set_paragraph_border(header_p, color="B7DCD5", size="8")

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_p.paragraph_format.space_before = Pt(0)
    footer_p.paragraph_format.space_after = Pt(0)
    footer_p.add_run("AI GURU  •  ")
    add_page_number(footer_p)
    for run in footer_p.runs:
        if not run._r.xpath(".//w:fldChar"):
            set_run_font(run, size=9, color=MUTED)


def add_hyperlink(paragraph, label: str, url: str, *, size: float = 10.5) -> None:
    """Add a clickable external hyperlink while keeping the document readable."""
    relationship_id = paragraph.part.relate_to(
        url,
        RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)
    properties.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "176B6A")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    size_element = OxmlElement("w:sz")
    size_element.set(qn("w:val"), str(round(size * 2)))
    properties.append(size_element)
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = label
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline_text(paragraph, text: str, *, default_size: float = 11) -> None:
    """Add a small, forgiving subset of Markdown inline formatting."""
    pattern = re.compile(
        r"(\[[^\]]+\]\([^)]*\)|\*\*.*?\*\*|`.*?`|\*.*?\*)"
    )
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size=default_size, color="1F2937")
        token = match.group(0)
        link_match = re.fullmatch(r"\[([^\]]+)\]\(([^)]*)\)", token)
        if link_match:
            add_hyperlink(
                paragraph,
                link_match.group(1),
                link_match.group(2),
                size=default_size,
            )
        elif token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=default_size, color=INK, bold=True)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name=MONO_FONT, size=max(default_size - 1, 8.5), color=INK)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=default_size, color="1F2937", italic=True)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=default_size, color="1F2937")


def clean_table_cell(value: str) -> str:
    # Keep Markdown links intact so table cells can become real Word hyperlinks.
    return value.strip()


def parse_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        if not line.strip().startswith("|"):
            continue
        raw = line.strip().strip("|")
        cells = [clean_table_cell(part) for part in raw.split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)
    return rows


def table_widths(column_count: int) -> list[int]:
    presets = {
        1: [9360],
        2: [3000, 6360],
        3: [1650, 3100, 4610],
        4: [1300, 2500, 2600, 2960],
        5: [1150, 1800, 1950, 2050, 2410],
    }
    if column_count in presets:
        return presets[column_count]
    base = CONTENT_WIDTH_DXA // column_count
    widths = [base] * column_count
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def add_data_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    rows = [row + [""] * (column_count - len(row)) for row in rows]
    table = doc.add_table(rows=len(rows), cols=column_count)
    set_table_geometry(table, table_widths(column_count))
    table.style = "Table Grid"
    mark_header_row(table.rows[0])

    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = ""
            if row_index == 0:
                fill = LIGHT_TEAL
            elif row_index % 2 == 0:
                fill = "F8FBFB"
            else:
                fill = WHITE
            set_cell_shading(cell, fill)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            add_inline_text(p, value, default_size=9.3)
            if row_index == 0:
                for run in p.runs:
                    run.bold = True
                    set_run_font(run, size=9.3, color=NAVY, bold=True)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def add_callout(doc: Document, text_lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA], indent=TABLE_INDENT_DXA)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "FFF7F0")
    set_cell_border(cell, color="F1CDB5", size="8")
    cell.text = ""
    for index, line in enumerate(text_lines):
        p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        p.style = doc.styles["Narrative Body"]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5 if index < len(text_lines) - 1 else 0)
        add_inline_text(p, line, default_size=10.5)
        for run in p.runs:
            set_run_font(run, size=10.5, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code_block(doc: Document, code_lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA], indent=TABLE_INDENT_DXA)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F1F4F5")
    set_cell_border(cell, color="C8D2D8", size="6")
    cell.text = ""
    text = "\n".join(code_lines)
    p = cell.paragraphs[0]
    p.style = doc.styles["Code Block"]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_together = True
    p.paragraph_format.line_spacing = 1.0
    for index, line in enumerate(text.splitlines()):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, name=MONO_FONT, size=8.4, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_flow_strip(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=5)
    widths = [1750, 1860, 1860, 1860, 2030]
    set_table_geometry(table, widths)
    labels = ["CÂU HỎI", "DỮ LIỆU", "PHÉP TÍNH", "KIỂM TRA", "CÂU TRẢ LỜI"]
    fills = ["F0F3F5", "E5F2EF", "FFFFFF", "FEF0E7", "0B2545"]
    for index, (cell, label, fill) in enumerate(zip(table.rows[0].cells, labels, fills)):
        cell.text = ""
        set_cell_shading(cell, fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(label)
        set_run_font(run, size=9, color=INK, bold=True)
        if index == len(labels) - 1:
            for run in p.runs:
                set_run_font(run, size=9, color=WHITE, bold=True)
        elif index < len(labels) - 1:
            run = p.add_run("  →")
            set_run_font(run, size=9, color=TEAL, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_cover(doc: Document) -> None:
    for _ in range(2):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    run = kicker.add_run("PRODUCT THESIS  •  SYSTEM DESCRIPTION  •  V2")
    set_run_font(run, size=10, color=TEAL, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("AI GURU")
    set_run_font(run, size=33, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    subtitle.paragraph_format.keep_with_next = True
    run = subtitle.add_run("Từ câu hỏi tài chính đến câu trả lời có thể kiểm chứng")
    set_run_font(run, size=16, color=SLATE)

    description = doc.add_paragraph()
    description.alignment = WD_ALIGN_PARAGRAPH.CENTER
    description.paragraph_format.space_after = Pt(26)
    run = description.add_run("Thuyết minh sản phẩm, dữ liệu, kiến trúc, pipeline và hướng dẫn vận hành")
    set_run_font(run, size=10.5, color=MUTED, italic=True)

    add_image(
        doc,
        IMAGE_DIR / "cover_flow_v2.png",
        alt="Mạch sản phẩm AI GURU từ câu hỏi đến câu trả lời có thể kiểm chứng",
        width=6.42,
        caption=None,
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    metadata = [
        ("Sản phẩm", "AI GURU / ViFinQA"),
        ("Bài toán", "Hỏi đáp trên báo cáo tài chính"),
        ("Phiên bản", "2.0 — bản trực quan"),
        ("Ngày cập nhật", "31/08/2026"),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    set_table_geometry(table, [2400, 6960])
    for row, (label, value) in zip(table.rows, metadata):
        for cell in row.cells:
            set_cell_border(cell, color="E5E7EB", size="4")
        row.cells[0].text = ""
        row.cells[1].text = ""
        set_cell_shading(row.cells[0], LIGHT_GRAY)
        p = row.cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        set_run_font(r, size=9.5, color=INK, bold=True)
        p = row.cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_run_font(r, size=9.5, color="374151")

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(18)
    run = note.add_run("Một câu trả lời tốt cần đúng nguồn, đúng ngữ nghĩa, đúng phép tính và có thể truy vết.")
    set_run_font(run, size=10, color=CORAL_DARK, italic=True)
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int) -> None:
    if level == 1:
        chapter = re.match(r"^(\d+)\.\s+", text)
        if chapter:
            kicker = doc.add_paragraph(style="Chapter Kicker")
            kicker.add_run(f"CHAPTER {int(chapter.group(1)):02d}")
            for run in kicker.runs:
                set_run_font(run, size=8.5, color=TEAL, bold=True)
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    if level == 1:
        set_paragraph_border(p, color="B7DCD5", size="8")


def add_image(
    doc: Document,
    image_path: Path,
    *,
    alt: str,
    width: float = 6.42,
    caption: str | None = None,
) -> None:
    if not image_path.is_file():
        raise FileNotFoundError(f"Diagram PNG not found: {image_path}")
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.keep_with_next = bool(caption)
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(image_path), width=Inches(width))
    inline_shape._inline.docPr.set("descr", alt)
    inline_shape._inline.docPr.set("title", alt)
    if caption:
        caption_p = doc.add_paragraph(style="Image Caption")
        caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_p.paragraph_format.keep_with_next = False
        run = caption_p.add_run(caption)
        set_run_font(run, size=9.5, color=MUTED, italic=True)


def add_paragraph(doc: Document, text: str) -> None:
    if not text.strip():
        return
    p = doc.add_paragraph(style="Narrative Body")
    add_inline_text(p, text)


def next_numbering_id(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    return max(ids, default=0) + 1


def next_abstract_numbering_id(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    return max(ids, default=0) + 1


def create_numbering(doc: Document, *, ordered: bool) -> int:
    """Create a real Word numbering definition and return its numId."""
    numbering = doc.part.numbering_part.element
    abstract_id = next_abstract_numbering_id(doc)
    num_id = next_numbering_id(doc)

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")

    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if ordered else "•")
    level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "280")
    p_pr.append(ind)
    level.append(p_pr)
    if not ordered:
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), "Symbol")
        r_fonts.set(qn("w:hAnsi"), "Symbol")
        r_pr.append(r_fonts)
        level.append(r_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def add_list(doc: Document, items: list[str], *, ordered: bool) -> None:
    if ordered:
        num_id = create_numbering(doc, ordered=True)
    else:
        num_id = getattr(doc, "_aiguru_bullet_num_id", None)
        if num_id is None:
            num_id = create_numbering(doc, ordered=False)
            doc._aiguru_bullet_num_id = num_id
    for item in items:
        p = doc.add_paragraph(style="Narrative Body")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.194)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.208
        apply_numbering(p, num_id)
        add_inline_text(p, item)


def add_markdown_body(doc: Document, lines: list[str]) -> None:
    index = 0
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            add_paragraph(doc, " ".join(part.strip() for part in paragraph_buffer))
            paragraph_buffer = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading_match:
            flush_paragraph()
            add_heading(doc, heading_match.group(2).strip(), len(heading_match.group(1)))
            index += 1
            continue

        if stripped == "---":
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(8)
            set_paragraph_border(p, color="D7DEE8", size="4")
            index += 1
            continue

        image_match = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
        if image_match:
            flush_paragraph()
            alt_text = image_match.group(1).strip() or "AI GURU diagram"
            image_ref = Path(image_match.group(2).strip())
            image_key = image_ref.stem
            image_path = IMAGE_DIR / f"{image_key}.png"
            add_image(
                doc,
                image_path,
                alt=alt_text,
                width=6.42,
                caption=IMAGE_CAPTIONS.get(image_key),
            )
            index += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            callout_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"): 
                value = lines[index].strip()[1:].strip()
                if value:
                    callout_lines.append(value)
                index += 1
            if callout_lines:
                add_callout(doc, callout_lines)
            continue

        if stripped.startswith("```"):
            flush_paragraph()
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index].rstrip())
                index += 1
            if index < len(lines):
                index += 1
            add_code_block(doc, code_lines)
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            flush_paragraph()
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_data_table(doc, parse_table(table_lines))
            continue

        bullet_match = re.match(r"^[-*]\s+(.+)$", stripped)
        numbered_match = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if bullet_match or numbered_match:
            flush_paragraph()
            items: list[str] = []
            ordered = numbered_match is not None
            while index < len(lines):
                current = lines[index].strip()
                match = re.match(r"^\d+[.)]\s+(.+)$", current) if ordered else re.match(r"^[-*]\s+(.+)$", current)
                if not match:
                    break
                items.append(match.group(1))
                index += 1
            add_list(doc, items, ordered=ordered)
            continue

        paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()


def build(output_path: Path) -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    separator = next((i for i, line in enumerate(lines) if line.strip() == "---"), None)
    if separator is None:
        raise ValueError("Markdown source has no front-matter separator")

    doc = Document()
    configure_styles(doc)
    configure_page_furniture(doc)
    add_cover(doc)
    add_markdown_body(doc, lines[separator + 1 :])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


if __name__ == "__main__":
    target = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_OUTPUT
    build(target)
    print(target)
